import os
import sys
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

def build_presentation():
    prs = Presentation()
    # Set 16:9 widescreen layout
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    
    # Blank slide layout
    blank_layout = prs.slide_layouts[6]
    
    # Color Palette (Dark Tech Theme)
    BG_DARK = RGBColor(15, 23, 42)       # #0F172A Slate 900
    CARD_BG = RGBColor(30, 41, 59)      # #1E293B Slate 800
    CARD_BORDER = RGBColor(51, 65, 85)  # #334155 Slate 700
    TEXT_MAIN = RGBColor(248, 250, 252) # #F8FAFC White/Slate 50
    TEXT_MUTED = RGBColor(148, 163, 184)# #94A3B8 Slate 400
    ACCENT_BLUE = RGBColor(59, 130, 246)# #3B82F6 Primary Blue
    ACCENT_GREEN = RGBColor(16, 185, 129)# #10B981 Emerald Green
    ACCENT_PURPLE = RGBColor(139, 92, 246)# #8B5CF6 Purple
    ACCENT_AMBER = RGBColor(245, 158, 11)# #F59E0B Amber/Gold

    def add_bg(slide):
        bg = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, 0, 0, Inches(13.333), Inches(7.5))
        bg.fill.solid()
        bg.fill.fore_color.rgb = BG_DARK
        bg.line.fill.background() # No line

    def add_header(slide, title_text, subtitle_text=""):
        tb = slide.shapes.add_textbox(Inches(0.8), Inches(0.4), Inches(11.733), Inches(1.1))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
        
        p = tf.paragraphs[0]
        p.text = title_text
        p.font.name = "Arial"
        p.font.size = Pt(26)
        p.font.bold = True
        p.font.color.rgb = TEXT_MAIN
        
        if subtitle_text:
            p2 = tf.add_paragraph()
            p2.text = subtitle_text
            p2.font.name = "Arial"
            p2.font.size = Pt(14)
            p2.font.color.rgb = ACCENT_BLUE
            p2.space_before = Pt(4)

    def add_card(slide, left, top, width, height, title, items, border_color=CARD_BORDER, header_color=ACCENT_BLUE):
        # Card shape background
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(width), Inches(height))
        card.fill.solid()
        card.fill.fore_color.rgb = CARD_BG
        card.line.color.rgb = border_color
        card.line.width = Pt(1.5)

        # Text Frame
        tb = slide.shapes.add_textbox(Inches(left + 0.25), Inches(top + 0.25), Inches(width - 0.5), Inches(height - 0.5))
        tf = tb.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0

        p = tf.paragraphs[0]
        p.text = title
        p.font.name = "Arial"
        p.font.size = Pt(18)
        p.font.bold = True
        p.font.color.rgb = header_color

        for item in items:
            p = tf.add_paragraph()
            p.space_before = Pt(8)
            p.font.name = "Arial"
            p.font.size = Pt(13)
            p.font.color.rgb = TEXT_MAIN
            
            if isinstance(item, tuple):
                # Bullet title + description
                run1 = p.add_run()
                run1.text = "• " + item[0] + ": "
                run1.font.bold = True
                run1.font.color.rgb = ACCENT_BLUE
                
                run2 = p.add_run()
                run2.text = item[1]
                run2.font.color.rgb = TEXT_MAIN
            else:
                run = p.add_run()
                run.text = "• " + str(item)
                run.font.color.rgb = TEXT_MAIN

    # ==========================================
    # SLIDE 1: COVER
    # ==========================================
    slide1 = prs.slides.add_slide(blank_layout)
    add_bg(slide1)
    
    # Hero container
    hero = slide1.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(1.0), Inches(1.2), Inches(11.333), Inches(5.1))
    hero.fill.solid()
    hero.fill.fore_color.rgb = CARD_BG
    hero.line.color.rgb = ACCENT_BLUE
    hero.line.width = Pt(2)

    tb = slide1.shapes.add_textbox(Inches(1.4), Inches(1.6), Inches(10.533), Inches(4.3))
    tf = tb.text_frame
    tf.word_wrap = True

    p0 = tf.paragraphs[0]
    p0.text = "HRM AI"
    p0.font.name = "Arial"
    p0.font.size = Pt(44)
    p0.font.bold = True
    p0.font.color.rgb = ACCENT_BLUE

    p1 = tf.add_paragraph()
    p1.text = "Hệ thống Tuyển dụng & Quản trị Nhân sự Tích hợp AI"
    p1.font.name = "Arial"
    p1.font.size = Pt(24)
    p1.font.bold = True
    p1.font.color.rgb = TEXT_MAIN
    p1.space_before = Pt(8)

    p2 = tf.add_paragraph()
    p2.text = "Từ nhu cầu tuyển dụng đến ngày đầu tiên của nhân viên — Theo dõi thông minh qua Telegram"
    p2.font.name = "Arial"
    p2.font.size = Pt(16)
    p2.font.color.rgb = ACCENT_GREEN
    p2.space_before = Pt(8)

    p3 = tf.add_paragraph()
    p3.text = "Quy trình khép kín: Nhu cầu → Đăng tin → AI Sàng lọc CV → Phỏng vấn → Offer → Pre-boarding → Xác nhận Đi làm → Telegram Management"
    p3.font.name = "Arial"
    p3.font.size = Pt(13)
    p3.font.color.rgb = TEXT_MUTED
    p3.space_before = Pt(24)

    p4 = tf.add_paragraph()
    p4.text = "Đơn vị: HRM AI Project Team | Công nghệ: React 19, Spring Boot, Google Gemini AI, Telegram Bot"
    p4.font.name = "Arial"
    p4.font.size = Pt(12)
    p4.font.color.rgb = ACCENT_PURPLE
    p4.space_before = Pt(16)

    # ==========================================
    # SLIDE 2: BÀI TOÁN & GIẢI PHÁP
    # ==========================================
    slide2 = prs.slides.add_slide(blank_layout)
    add_bg(slide2)
    add_header(slide2, "Slide 2 — Bài toán thực tế & Giải pháp HRM AI", "Chuyển đổi quy trình thủ công đứt gãy thành Nền tảng hợp nhất Nhanh - Đúng - Minh bạch")

    add_card(slide2, 0.8, 1.6, 5.7, 5.3, "Thách thức Quy trình Truyền thống", [
        ("Nhập liệu thủ công", "CV gửi qua nhiều kênh, HR mất nhiều giờ đọc và nhập dữ liệu."),
        ("Tiêu chí thiếu nhất quán", "Đánh giá ứng viên không đồng bộ giữa HR và bộ phận chuyên môn."),
        ("Tuyển vượt chỉ tiêu", "Thiếu cơ chế Seat Ledger giữ suất real-time khi phát hành Offer."),
        ("Đứt gãy dữ liệu", "Hồ sơ ứng viên trúng tuyển không tự động nối với hồ sơ Nhân viên."),
        ("Lãnh đạo thiếu công cụ", "Phụ thuộc báo cáo chậm, khó tra cứu & duyệt nhanh khi di chuyển.")
    ], border_color=ACCENT_AMBER, header_color=ACCENT_AMBER)

    add_card(slide2, 6.8, 1.6, 5.7, 5.3, "Giải pháp HRM AI Nổi bật", [
        ("Nhanh hơn", "Gemini AI tự động cấu trúc hóa CV & tính Fit Score theo bằng chứng trong vài giây."),
        ("Đúng hơn", "Kiểm soát qua State Machine nghiêm ngặt, tự động giữ suất tuyển (Seat Ledger)."),
        ("Minh bạch hơn", "Audit Log đầy đủ, cờ cảnh báo sao chép JD, người có thẩm quyền quyết định cuối."),
        ("Nối liền mạch", "Tự động chuyển Offer Accepted → Hồ sơ Nhân viên (Pre-boarding) → Kích hoạt tài khoản."),
        ("Lãnh đạo đồng hành", "Telegram Bot tra cứu dữ liệu real-time và phê duyệt tức thì.")
    ], border_color=ACCENT_GREEN, header_color=ACCENT_GREEN)

    # ==========================================
    # SLIDE 3: LUỒNG TỔNG THỂ (12 BƯỚC)
    # ==========================================
    slide3 = prs.slides.add_slide(blank_layout)
    add_bg(slide3)
    add_header(slide3, "Slide 3 — Luồng Tuyển dụng & Tiếp nhận Nhân sự Tổng thể", "Quy trình 12 bước xuyên suốt từ Nhu cầu tuyển dụng đến Ngày đi làm đầu tiên của Nhân viên")

    col_w = 3.65
    gap_x = 0.35
    row_h = 1.15
    gap_y = 0.15
    
    steps_data = [
        ("1. Tạo nhu cầu tuyển dụng", "Trưởng phòng lập Job Requisition (DRAFT)"),
        ("2. Phê duyệt Requisition", "CEO/Quản lý duyệt → Cấp Hiring Seats (AVAILABLE)"),
        ("3. Đăng tin tuyển dụng", "HR tạo Posting, khóa cặp Version Tiêu chí & AI Profile"),
        ("4. Ứng viên nộp CV", "Nộp trực tuyến qua Public Form (Không thu CCCD)"),
        ("5. AI Phân tích CV", "PDFBox extract + Gemini AI chấm Fit Score theo bằng chứng"),
        ("6. HR & Chuyên môn duyệt", "Sàng lọc CV qua 2 vòng (PENDING_HR → PENDING_TECH)"),
        ("7. Phỏng vấn 2 Vòng", "Vòng 1 Chuyên môn & Vòng 2 Quản lý/HR đánh giá"),
        ("8. Duyệt & Gửi Offer", "Lập Offer, duyệt ngân sách, gửi ứng viên & GIỮ SUẤT (RESERVED)"),
        ("9. Ứng viên Chấp nhận", "Ứng viên click OFFER_ACCEPTED, kích hoạt bridge event"),
        ("10. Pre-boarding & Checklist", "Tạo Employee (PRE_BOARDING), checklist CCCD, account DISABLED"),
        ("11. HR Xác nhận Đi làm", "Ngày đầu đi làm: HR bấm xác nhận JOINED"),
        ("12. Kích hoạt Nhân viên", "Employee → EMPLOYED, Seat → JOINED, Account → ACTIVE")
    ]

    for idx, (stitle, sdesc) in enumerate(steps_data):
        r = idx // 3
        c = idx % 3
        left = 0.8 + c * (col_w + gap_x)
        top = 1.6 + r * (row_h + gap_y)
        
        box = slide3.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(left), Inches(top), Inches(col_w), Inches(row_h))
        box.fill.solid()
        box.fill.fore_color.rgb = CARD_BG
        box.line.color.rgb = ACCENT_BLUE if idx < 8 else ACCENT_GREEN
        box.line.width = Pt(1)

        tf = box.text_frame
        tf.word_wrap = True
        tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = Inches(0.1)

        p = tf.paragraphs[0]
        p.text = stitle
        p.font.name = "Arial"
        p.font.size = Pt(12)
        p.font.bold = True
        p.font.color.rgb = ACCENT_BLUE if idx < 8 else ACCENT_GREEN

        p2 = tf.add_paragraph()
        p2.text = sdesc
        p2.font.name = "Arial"
        p2.font.size = Pt(10)
        p2.font.color.rgb = TEXT_MAIN
        p2.space_before = Pt(2)

    # ==========================================
    # SLIDE 4: KHỞI TẠO NHU CẦU & STRATEGY
    # ==========================================
    slide4 = prs.slides.add_slide(blank_layout)
    add_bg(slide4)
    add_header(slide4, "Slide 4 — Khởi tạo Nhu cầu & Đăng tin Tuyển dụng", "Quản lý Requisition, Sổ kế toán Suất tuyển (Seat Ledger) và Khóa tiêu chí AI")

    add_card(slide4, 0.8, 1.6, 3.7, 5.3, "1. Yêu cầu Tuyển dụng", [
        ("Khởi tạo Requisition", "Trưởng phòng lập yêu cầu theo vị trí, phòng ban, headcount, khung lương."),
        ("Duyệt đa cấp", "Phê duyệt theo ma trận phân quyền backend (DRAFT → APPROVED)."),
        ("Ràng buộc ngân sách", "Ngăn chặn tuyển dụng ngoài kế hoạch hoặc vượt định mức lương.")
    ], border_color=ACCENT_BLUE, header_color=ACCENT_BLUE)

    add_card(slide4, 4.8, 1.6, 3.7, 5.3, "2. Quản lý Suất tuyển (Seat Ledger)", [
        ("Cấp Suất tự động", "Khi Requisition APPROVED, tự động sinh các bản ghi hiring_seats ở AVAILABLE."),
        ("Ràng buộc nghiêm ngặt", "Mọi đợt gửi Offer đều phải gắn với 1 Suất tuyển cụ thể."),
        ("Chống overbook", "Không bao giờ xảy ra tình trạng ứng viên chấp nhận offer nhưng báo hết chỉ tiêu.")
    ], border_color=ACCENT_GREEN, header_color=ACCENT_GREEN)

    add_card(slide4, 8.8, 1.6, 3.7, 5.3, "3. Đóng băng Tiêu chí AI", [
        ("Posting Open", "HR tạo Posting & phải xác nhận bộ tiêu chí chấm và AI Scoring Profile."),
        ("Version bất biến", "Khóa cặp (criteria_version, scoring_profile_version) bất biến khi đăng tin."),
        ("Audit re-run", "Khi sửa tiêu chí, chỉ chạy lại AI cho các hồ sơ đang ở giai đoạn đánh giá.")
    ], border_color=ACCENT_PURPLE, header_color=ACCENT_PURPLE)

    # ==========================================
    # SLIDE 5: AI PHÂN TÍCH & CHẤM CV
    # ==========================================
    slide5 = prs.slides.add_slide(blank_layout)
    add_bg(slide5)
    add_header(slide5, "Slide 5 — AI Phân tích, Sàng lọc & Chấm CV theo Bằng chứng", "Bảo mật tải lên, Trích xuất PDFBox, Gemini AI Fit Score & Cảnh báo bất thường")

    add_card(slide5, 0.8, 1.6, 5.7, 5.3, "Quy trình Xử lý CV & AI Gemini", [
        ("Nộp Form trực tuyến", "Ứng viên điền thông tin & nộp CV (PDF/DOCX). Tuyệt đối KHÔNG thu thập CCCD."),
        ("Kiểm tra An toàn File", "Xác minh mime-type, dung lượng, số trang, chặn virus/macro và file mã hóa."),
        ("PDFBox Text Extraction", "Trích xuất toàn bộ văn bản thô từ tệp PDF CV chuẩn xác."),
        ("Gemini AI Parsing", "Phân tích cấu trúc: Học vấn, Kinh nghiệm, Kỹ năng & Thành tựu thực tế.")
    ], border_color=ACCENT_BLUE, header_color=ACCENT_BLUE)

    add_card(slide5, 6.8, 1.6, 5.7, 5.3, "Fit Score theo Bằng chứng & Safety", [
        ("Evidence-Based Scoring", "Chấm theo BẰNG CHỨNG thực tế thay vì từ khóa suông. Phân biệt Khai báo vs Thực chứng."),
        ("AI Fraud Flag", "Phát hiện & cảnh báo ứng viên sao chép nguyên văn mô tả công việc (JD Mirroring)."),
        ("Interview Questions", "AI tự động sinh danh sách câu hỏi phỏng vấn xoáy sâu vào các điểm cần xác minh."),
        ("Nguyên tắc bất biến", "AI chỉ gợi ý và sắp xếp danh sách. HR & Chuyên môn luôn giữ quyền quyết định!")
    ], border_color=ACCENT_AMBER, header_color=ACCENT_AMBER)

    # ==========================================
    # SLIDE 6: PHỎNG VẤN & PHÊ DUYỆT OFFER
    # ==========================================
    slide6 = prs.slides.add_slide(blank_layout)
    add_bg(slide6)
    add_header(slide6, "Slide 6 — Phỏng vấn có kiểm soát & Phê duyệt Offer an toàn", "Quy trình phỏng vấn 2 vòng, Duyệt Offer nội bộ và Cơ chế Giữ suất tuyển (Seat Reservation)")

    add_card(slide6, 0.8, 1.6, 5.7, 5.3, "Quy trình Phỏng vấn 2 Vòng", [
        ("Duyệt CV 2 cấp", "HR duyệt CV cơ bản (PENDING_HR_CV_REVIEW) → Phòng chuyên môn duyệt (PENDING_TECH_CV_REVIEW)."),
        ("Phỏng vấn Vòng 1", "Đánh giá năng lực chuyên môn thực hành & giải quyết bài toán."),
        ("Phỏng vấn Vòng 2", "Đánh giá sự phù hợp văn hóa & định hướng phát triển từ Quản lý/HR."),
        ("Hỗ trợ Revert", "Cho phép trả hồ sơ về bước trước có lý do kèm theo nếu cần đánh giá lại.")
    ], border_color=ACCENT_BLUE, header_color=ACCENT_BLUE)

    add_card(slide6, 6.8, 1.6, 5.7, 5.3, "Tạo & Phê duyệt Offer An toàn", [
        ("Tách biệt Offer & Dispatch", "Tách điều khoản Offer (bất biến) và Lần gửi Dispatch (token, deadline)."),
        ("Phê duyệt Nội bộ", "Duyệt theo chính sách ngân sách Requisition. Cảnh báo cờ vượt khung lương."),
        ("Giữ Suất tuyển (Seat Reserved)", "Nối từ OFFER_SENT: Suất tuyển chuyển từ AVAILABLE → RESERVED."),
        ("Link Phản hồi Bảo mật", "Ứng viên nhận link bảo mật có token phản hồi (Accept / Decline / Negotiate).")
    ], border_color=ACCENT_GREEN, header_color=ACCENT_GREEN)

    # ==========================================
    # SLIDE 7: TỪ OFFER CHẤP NHẬN ĐẾN NHÂN VIÊN
    # ==========================================
    slide7 = prs.slides.add_slide(blank_layout)
    add_bg(slide7)
    add_header(slide7, "Slide 7 — Từ Offer Chấp nhận đến Nhân viên Chính thức", "Cầu nối Tuyển dụng - Nhân sự (Idempotent Bridge), Pre-boarding, Cấp tài khoản và JOINED")

    add_card(slide7, 0.8, 1.6, 5.7, 5.3, "Cầu nối Idempotent & Pre-boarding", [
        ("Event OFFER_ACCEPTED", "Ứng viên bấm chấp nhận offer → Phát event Outbox với conversion_key chống trùng."),
        ("Tạo Person & Employee", "Khởi tạo Person (PROVISIONAL) & Employee ở trạng thái PRE_BOARDING."),
        ("Checklist & Hồ sơ", "Tự động sinh Pre-boarding Checklist: Thu thập CCCD, Bằng cấp, Ngân hàng."),
        ("Hợp đồng & Lương ban đầu", "Tạo employment_record & compensation_record ở trạng thái SCHEDULED.")
    ], border_color=ACCENT_PURPLE, header_color=ACCENT_PURPLE)

    add_card(slide7, 6.8, 1.6, 5.7, 5.3, "Tài khoản DISABLED & Xác nhận JOINED", [
        ("Tài khoản An toàn", "Gửi yêu cầu tạo tài khoản hệ thống. Tạo trước ngày đi làm ở trạng thái DISABLED."),
        ("Xác nhận Đi làm (JOINED)", "Vào ngày nhận việc, HR bấm 'Xác nhận đã đi làm' trên hệ thống."),
        ("Kích hoạt Đồng bộ", "Employee → EMPLOYED, Hiring Seat → JOINED, Account → ACTIVE."),
        ("Hủy nhận việc", "Nếu ứng viên không đến → ONBOARD_CANCELLED, hủy tài khoản & nhả lại Suất tuyển.")
    ], border_color=ACCENT_GREEN, header_color=ACCENT_GREEN)

    # ==========================================
    # SLIDE 8: TELEGRAM BOT FOR MANAGEMENT
    # ==========================================
    slide8 = prs.slides.add_slide(blank_layout)
    add_bg(slide8)
    add_header(slide8, "Slide 8 — Telegram HRM AI Bot: Điều hành & Phê duyệt Tức thì", "Trợ lý thông minh dành cho CEO & Giám đốc — Tra cứu dữ liệu real-time và Phê duyệt nhanh 24/7")

    add_card(slide8, 0.8, 1.6, 5.7, 5.3, "Bộ Lệnh Tra cứu & Duyệt nhanh", [
        ("/dashboard", "Xem tổng quan chỉ số nhân sự toàn công ty real-time."),
        ("/nhanvien & /tuyendung", "Thống kê quy mô nhân sự, phòng ban & đợt tuyển dụng."),
        ("/chamcong & /luong", "Kiểm tra tình hình đi làm trong ngày & tổng chi phí lương."),
        ("/request, /duyet, /tuchoi", "Duyệt hoặc từ chối yêu cầu tuyển dụng, offer theo ID tức thì.")
    ], border_color=ACCENT_BLUE, header_color=ACCENT_BLUE)

    add_card(slide8, 6.8, 1.6, 5.7, 5.3, "Hỏi đáp Tiếng Việt & Keyword Fallback", [
        ("Truy vấn Tự nhiên", "Hỏi câu hỏi tiếng Việt: 'Hiện công ty có bao nhiêu nhân viên?', 'Hôm nay ai đi muộn?'"),
        ("Gemini Intent Recognition", "Gemini AI phân tích ý định câu hỏi → gọi API HRM lấy dữ liệu thực."),
        ("Cơ chế Fallback An toàn", "Nếu kết nối AI sự cố, Bot tự chuyển sang nhận diện Keyword phục vụ liên tục."),
        ("Bảo mật Phân quyền", "Bot xác thực theo Chat ID và Số điện thoại chỉ dành cho Quản lý được cấp quyền.")
    ], border_color=ACCENT_GREEN, header_color=ACCENT_GREEN)

    # ==========================================
    # SLIDE 9: KIẾN TRÚC CÔNG NGHỆ & SAFEGUARD
    # ==========================================
    slide9 = prs.slides.add_slide(blank_layout)
    add_bg(slide9)
    add_header(slide9, "Slide 9 — Kiến trúc Công nghệ & Cấu hình Động (No-Hardcode)", "Hệ thống chuẩn Enterprise: React 19, Spring Boot, MySQL Flyway, Outbox Pattern & Dynamic Config")

    add_card(slide9, 0.8, 1.6, 5.7, 5.3, "Kiến trúc Multi-Tier Modern", [
        ("Frontend Modern", "React 19 + Vite + Tailwind CSS + Recharts + WebSocket (STOMP) real-time notification."),
        ("Backend Enterprise", "Java 17 + Spring Boot + Spring Security (JWT) + Multi-step Approval Engine."),
        ("Data & Event Storage", "MySQL + Flyway DB Migration + Outbox Event Pattern có retry tự động."),
        ("AI & File Engine", "Google Gemini API + Apache PDFBox xử lý trích xuất văn bản CV.")
    ], border_color=ACCENT_BLUE, header_color=ACCENT_BLUE)

    add_card(slide9, 6.8, 1.6, 5.7, 5.3, "Nguyên tắc 'Không Hardcode' (No-Hardcode)", [
        ("Database-Driven Config", "Toàn bộ tham số nghiệp vụ nằm trong system_configurations & scoring_parameters."),
        ("Thay đổi tức thì", "HR Head/Admin sửa hệ số điểm AI, SLA, dung lượng file trên UI → Hệ thống đổi ngay không cần redeploy."),
        ("Version bất biến", "Profile chấm điểm AI được đánh version bất biến, đảm bảo tính nhất quán audit."),
        ("Bảo mật & Phân quyền", "Phân quyền chi tiết (RBAC), Audit log append-only cho mọi thao tác chuyển trạng thái.")
    ], border_color=ACCENT_PURPLE, header_color=ACCENT_PURPLE)

    # ==========================================
    # SLIDE 10: GIÁ TRỊ DOANH NGHIỆP & DEMO
    # ==========================================
    slide10 = prs.slides.add_slide(blank_layout)
    add_bg(slide10)
    add_header(slide10, "Slide 10 — Giá trị Doanh nghiệp & Kịch bản Demo 3-5 Phút", "Tối ưu vận hành nhân sự toàn diện và Kịch bản trình diễn trực quan tính năng sản phẩm")

    add_card(slide10, 0.8, 1.6, 5.7, 5.3, "Giá trị mang lại cho Doanh nghiệp", [
        ("Cho HR", "Tự động hóa 80% thao tác đọc/nhập CV, loại bỏ bỏ sót hồ sơ, quản lý luồng tập trung."),
        ("Cho Quản lý", "Đánh giá ứng viên theo tiêu chí đồng nhất, nắm chính xác tiến độ từng hồ sơ."),
        ("Cho Lãnh đạo", "Kiểm soát ngân sách & chỉ tiêu (Seat Ledger), tra cứu & phê duyệt 24/7 qua Telegram."),
        ("Toàn vẹn Dữ liệu", "Nối liền mạch Tuyển dụng → Nhân viên, lưu vết Lịch sử (Audit Log) đầy đủ.")
    ], border_color=ACCENT_GREEN, header_color=ACCENT_GREEN)

    add_card(slide10, 6.8, 1.6, 5.7, 5.3, "Kịch bản Demo Thực tế (3 - 5 Phút)", [
        ("B1. Đăng tin & CV List", "Mở Tin tuyển dụng đang OPEN & danh sách ứng viên nộp CV."),
        ("B2. AI Fit Score Detail", "Xem Fit Score (chấm bằng chứng), Cảnh báo Fraud & Đề xuất câu hỏi phỏng vấn."),
        ("B3. Offer & Seat Reserve", "Chuyển phỏng vấn → Duyệt Offer → Suất tuyển chuyển RESERVED."),
        ("B4. Accept & Joined Flow", "Ứng viên click Accept → Hồ sơ vào Pre-boarding → HR bấm JOINED → Account ACTIVE."),
        ("B5. Telegram Live Query", "Mở Telegram hỏi 'Đang tuyển bao nhiêu vị trí?' & gõ /dashboard kết thúc.")
    ], border_color=ACCENT_AMBER, header_color=ACCENT_AMBER)

    output_path = os.path.join(os.path.dirname(__file__), "HRM_AI_Presentation.pptx")
    prs.save(output_path)
    print(f"Successfully generated PowerPoint presentation at: {output_path}")

if __name__ == "__main__":
    build_presentation()
